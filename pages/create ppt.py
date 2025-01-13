import openai
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
import requests
from io import BytesIO
import streamlit as st

# Initialize OpenAI client with your API key
openai_client = openai.OpenAI(api_key='') 

def query_vector_store(pdf_file, topic):
    assistant = openai_client.beta.assistants.create(
        name="PDF-to-Slide-Assistant",
        instructions="You are an expert at generating PowerPoint slides. Use your knowledge base to create concise slides on the topic provided.",
        model="gpt-3.5-turbo",
        tools=[{"type": "file_search"}],
    )
    
    # Use the file-like object directly (Streamlit's file uploader provides this)
    message_file = openai_client.files.create(
        file=pdf_file, purpose="assistants"
    )
    
    thread = openai_client.beta.threads.create(
        messages=[{
            "role": "user",
            "content": f'''
            Create a PowerPoint presentation on the topic '{topic}'. Use the following format for each slide:
            
            Slide [Number]:
            Title: [Title of the slide]
            Paragraph: A short paragraph (2-3 sentences) introducing the content of the slide.
            Primary Point 1: [First primary point]
            Secondary Point 1.1: [Supporting detail for Primary Point 1]
            Primary Point 2: [Second primary point]
            Secondary Point 2.1: [Supporting detail for Primary Point 2]
            
            Provide content for exactly 5 slides. Each slide should start with a paragraph, followed by two primary points and one secondary point for each primary point.
            Focus on making the content short, engaging, and easy to understand. Only include the content of the slides in the format provided, without any citations or extra content.
            ''',
            "attachments": [{"file_id": message_file.id, "tools": [{"type": "file_search"}]}],
        }]
    )
    
    run = openai_client.beta.threads.runs.create_and_poll(
        thread_id=thread.id, assistant_id=assistant.id
    )
    messages = list(openai_client.beta.threads.messages.list(thread_id=thread.id, run_id=run.id))
    message_content = messages[0].content[0].text
    return message_content.value

def fetch_image(query):
    response = openai_client.images.generate(
        model="dall-e-3",
        prompt=query,
        size="1024x1024",
        quality="standard",
        n=1,
    )
    image_url = response.data[0].url
    image_response = requests.get(image_url)
    if image_response.status_code == 200:
        return image_response.content  
    else:
        raise Exception(f"Failed to download the image. Status code: {image_response.status_code}")

def parse_content(content):
    slides = []
    current_slide = None
    current_primary = None

    for line in content.split('\n'):
        line = line.strip()

        if line.startswith('Slide'):
            if current_slide:
                slides.append(current_slide)
            current_slide = {'title': '', 'paragraph': '', 'points': []}
        
        elif line.startswith('Title:'):
            current_slide['title'] = line.split(':', 1)[-1].strip()

        elif line.startswith('Paragraph:'):
            current_slide['paragraph'] = line.split(':', 1)[-1].strip()

        elif line.startswith('Primary Point'):
            current_primary = {'text': line.split(':', 1)[-1].strip(), 'secondaries': []}
            current_slide['points'].append(current_primary)

        elif line.startswith('Secondary Point'):
            if current_primary:
                current_primary['secondaries'].append(line.split(':', 1)[-1].strip())

    if current_slide:
        slides.append(current_slide)
    
    return slides

def create_ppt(slides, topic):
    prs = Presentation()

    prs.slide_width = Inches(16)
    prs.slide_height = Inches(9)

    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)

    first_slide = prs.slides[0]
    title_shape = first_slide.shapes.title
    title_shape = first_slide.shapes.add_textbox(Inches(0), Inches(3.5), Inches(16), Inches(2))
    title_shape.text = topic
    title_shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    for paragraph in title_shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(48)

    for slide_content in slides:
        layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(layout)
        
        title_shape = slide.shapes.title
        title_shape.text = slide_content['title']
        title_shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        title_shape.top = Inches(0.5)
        title_shape.width = Inches(14)
        title_shape.height = Inches(1)

        for paragraph in title_shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(44)

        image_query = slide_content['title']
        image = fetch_image(image_query)
        image_stream = BytesIO(image)
        if image_stream:
            left = Inches(1)
            top = Inches(2)
            width = Inches(6)
            height = Inches(6)
            slide.shapes.add_picture(image_stream, left, top, width, height)

        content_left = Inches(8)
        content_top = Inches(2)
        content_width = Inches(7)
        content_height = Inches(6)

        content_box = slide.shapes.add_textbox(content_left, content_top, content_width, content_height)
        tf = content_box.text_frame
        tf.word_wrap = True

        p = tf.add_paragraph()
        p.text = slide_content['paragraph']
        p.level = 0
        p.font.size = Pt(24)
        p.line_spacing = 1.15

        for point in slide_content['points']:
            p = tf.add_paragraph()
            p.text = point['text']
            p.level = 0
            p.font.size = Pt(24)
            p.font.bold = True
            p.line_spacing = 1.15

            for secondary in point['secondaries']:
                sp = tf.add_paragraph()
                sp.text = '• ' + secondary
                sp.level = 1
                sp.font.size = Pt(18)
                sp.line_spacing = 1.15

    

    # Save the presentation to a BytesIO buffer
    pptx_buffer = BytesIO()
    prs.save(pptx_buffer)
    pptx_buffer.seek(0)  # Move the cursor to the start of the buffer
    return pptx_buffer

# Streamlit UI for file upload
def main():
    st.title("PDF to PowerPoint Generator")
    st.write("Upload a PDF file and specify a topic to generate a PowerPoint presentation.")

    # File uploader widget for PDF
    pdf_file = st.file_uploader("Upload PDF", type="pdf")

    # Text input for topic
    topic = st.text_input("Enter the topic")

    if pdf_file and topic:
        # Process the uploaded file and generate slides
        content = query_vector_store(pdf_file, topic)
        slides = parse_content(content)
        
        # Create PowerPoint in memory
        pptx_buffer = create_ppt(slides, topic)

        # Provide download button to the user
        st.success(f"Slides created successfully! Download the PPT from below:")
        st.download_button("Download PowerPoint", pptx_buffer, file_name=f"{topic.replace(' ', '_')}.pptx", mime="application/vnd.openxmlformats-officedocument.presentationml.presentation")

if __name__ == "__main__":
    main()
